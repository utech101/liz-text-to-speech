#Michael Jones
from tkinter import *
from tkinter import filedialog, messagebox
import pyttsx3
import PyPDF2
import docx
from newspaper import Article
from docx.opc.exceptions import PackageNotFoundError
from LexicalAnalyzer import LexicalAnalyzer
from Parser import Parser


# Converts text to speech based on the text passed to it
def textToSpeech(text):
    engine = pyttsx3.init()
    voices = engine.getProperty('voices')
    engine.setProperty('voice', voices[1].id)
    engine.setProperty('rate', 150)
    engine.say(text)
    engine.runAndWait()


# Displays the grammar analysis of text passed to it
def grammarAnalysis(text):
    LexicalAnalyzer.lexicalanalyzer(text)
    print("\n")
    Parser.parsing(text)


# Function used to convert text to speech, receive the text necessary for this conversion
# as well as error handling done on the text entered
def enterText():
    def submit():
        text = textField.get()
        if not text:
            messagebox.showerror("No Text Entered", "Please enter something for Liz to work.", parent=text_window)
        else:
            textToSpeech(text)
            grammarAnalysis(text)
            textField.delete(0, END)

    def back():
        text_window.destroy()

    def clear():
        textField.delete(0, END)

    text_window = Toplevel()
    text_window.title('TEXT')
    text_window.configure(bg="#FFFFFF")
    text_window.resizable(False, False)
    text_window.geometry("550x600+100+50")

    text_frame = LabelFrame(text_window, text="", padx=50, pady=50, borderwidth="4", bg="#FFFFFF")
    text_frame.grid(pady=100)

    label = Label(text_frame, text="Enter the desired text: ", bg="#FFFFFF")
    label.grid(row=0, column=0)
    label.configure(font=("Courier", 16))

    textField = Entry(text_frame)
    textField.grid(row=0, column=1, pady=(0, 50))

    submitBtn = Button(text_frame, text="SUBMIT", bg="#1A936F", fg="#FFFFFF", command=submit)
    submitBtn.grid(row=1, column=0)

    backBtn = Button(text_frame, text="BACK", bg="#1A936F", fg="#FFFFFF", command=back)
    backBtn.grid(row=1, column=1)

    clearBtn = Button(text_frame, text="CLEAR", bg="#1A936F", fg="#FFFFFF", command=clear)
    clearBtn.grid(row=1, column=2)


# Function used to convert the text present in files to speech, extract the information from the files necessary for this conversion
# as well as error handling done on the type of files and if the file is empty entered
def fileUpload():
    def upload():
        file_frame.pack()

        file_window.filename = filedialog.askopenfilename(initialdir=".", title="Select A File",
                                                          filetypes=[("Files", ".pdf .txt .docx")])

        if not file_window.filename:
            messagebox.showerror("No file selected", "Please upload a file", parent=file_window)
        elif not file_window.filename.lower().endswith(('.pdf', '.txt', '.docx')):
            messagebox.showerror("Invalid file extension",
                                 "Please upload a file with either a .pdf, .txt or .docx file extension",
                                 parent=file_window)
        elif file_window.filename.lower().endswith('.pdf'):
            reader = PyPDF2.PdfFileReader(file_window.filename)
            for i in range(reader.getNumPages()):
                page = reader.getPage(i)
                page_content = page.extractText()
                if page_content.isspace():
                    messagebox.showerror("Empty PDF File", "The PDF File selected is empty.", parent=file_window)
                    break
                else:
                    textToSpeech(page_content)
                    grammarAnalysis(page_content)
            reader.stream.close()
        elif file_window.filename.lower().endswith('.txt'):
            with open(file_window.filename, "r") as textFile:
                information = textFile.read()
                if not information:
                    messagebox.showerror("Empty Text File", "The Text File selected is empty.", parent=file_window)
                else:
                    textToSpeech(information)
                    grammarAnalysis(information)
        elif file_window.filename.lower().endswith('.docx'):
            try:
                doc = docx.Document(file_window.filename)
                completeText = '\n'.join([paragraph.text for paragraph in doc.paragraphs])
                textToSpeech(completeText)
                grammarAnalysis(completeText)
            except PackageNotFoundError:
                messagebox.showerror("Empty Word File", "The Word File selected is empty.", parent=file_window)

        file_frame.pack_forget()

    file_window = Toplevel()
    file_window.title('FILE UPLOAD')
    file_window.configure(bg="#FFFFFF")
    file_window.resizable(False, False)
    file_window.geometry("200x150+100+50")

    file_frame = LabelFrame(file_window, text="", padx=60, pady=60, borderwidth="4", bg="#FFFFFF")

    Button(file_window, text="SELECT FILE", bg="#1A936F", fg="#FFFFFF", command=upload).pack(pady=50)


# Function used to convert the text extracted from an article to speech
def onlineArticle():
    def readArticle():
        url = urlField.get()
        if not url:
            messagebox.showerror("No URL Entered", "Please enter something for Liz to work.", parent=article_window)
        else:
            article = Article(url, language="en")
            article.download()
            article.parse()
            title = "Title: " + article.title + "\n"
            author = "Author: " + str(article.authors) + "\n"
            published_date = "Published Date: " + str(article.publish_date) + "\n"
            article.nlp()
            summary = "Summary: " + article.summary + "\n"
            keywords = "Keywords: " + str(article.keywords) + "\n"
            articleInformation = title + author + published_date + summary + keywords + "\n"
            textToSpeech(articleInformation)
            grammarAnalysis(summary)
            urlField.delete(0, END)

    article_window = Toplevel()
    article_window.title('ARTICLE')
    article_window.configure(bg="#FFFFFF")
    article_window.resizable(False, False)
    article_window.geometry("570x170+100+50")

    article_frame = LabelFrame(article_window, text="", padx=60, pady=60, borderwidth="4", bg="#FFFFFF")
    article_frame.pack()

    articleLabel = Label(article_frame, text="Enter the article's URL: ", bg="#FFFFFF")
    articleLabel.configure(font=("Courier", 16))
    articleLabel.grid(row=0, column=0)

    urlField = Entry(article_frame)
    urlField.grid(row=0, column=1)

    readBtn = Button(article_frame, text="READ", bg="#1A936F", fg="#FFFFFF", padx=20, command=readArticle)
    readBtn.grid(row=1, column=0)


# Main Window of the Text to Speech Application that greets the user upon running the program
main_window = Tk()
main_window.title('LIZ TEXT TO SPEECH')
main_window.configure(bg="#FFFFFF")
main_window.resizable(False, False)
main_window.geometry("520x300+100+50")

main_frame = LabelFrame(main_window, text="", padx=60, pady=60, borderwidth="4", bg="#FFFFFF")
main_frame.grid(column=0, row=1, rowspan=2)

textLabel = Label(main_frame, text="LIZ TEXT TO SPEECH", bg="#FFFFFF", fg="#171A21")
textLabel.configure(font=("Courier", 17, "bold"))
textLabel.pack(pady=(0, 25))

textBtn = Button(main_frame, text="ENTER TEXT", bg="#3776ab", fg="#FFFFFF", pady=5, padx=5, command=enterText)
textBtn.pack(pady=(0, 50), side=TOP)
textBtn.place(anchor='ne', relx=0.2, rely=1)

articleBtn = Button(main_frame, text="READ AN ARTICLE", bg="#3776ab", fg="#FFFFFF", pady=5, padx=5, command=onlineArticle)
articleBtn.pack(pady=(0, 50), side=TOP)
articleBtn.place(anchor='nw', relx=0.8, rely=1)

uploadBtn = Button(main_frame, text="UPLOAD A FILE", bg="#3776ab", fg="#FFFFFF", pady=5, padx=5, command=fileUpload)
uploadBtn.pack(pady=(0, 50), side=TOP)
uploadBtn.place(anchor='n', relx=0.5, rely=1)



mainloop()
